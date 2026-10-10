from __future__ import annotations

import json
import types
from pathlib import Path

import pytest

from perf_analyzer.collector import (
    BUSINESS_HEADERS,
    FIXED_HEADERS,
    build_frame,
    build_system_info,
    map_metrics_snapshot,
    new_sheet_name,
    output_columns,
    read_metrics_snapshot,
    write_workbook,
)
from perf_analyzer.constants import PERFORMANCE_MODE
from perf_analyzer.io import read_single_workbook

SAMPLE_ROWS = [
    {
        "时间戳": f"2026-06-13 10:00:{second:02d}",
        "CPU%": 30.0 + second,
        "内存占用（MB）": 512.0 + second,
        "系统总内存": 16384,
        "内存占用百分比": round((512.0 + second) / 16384 * 100, 4),
        "线程": 40 + second % 5,
        "句柄": 800 + second,
        "io": round(1.2 + second / 10, 3),
        "磁盘读(MB/s)": round(0.4 + second / 20, 3),
        "磁盘写(MB/s)": round(0.2 + second / 40, 3),
        "IOPS": 40 + second,
        "网络上行(KB/s)": round(12.5 + second / 3, 2),
        "网络下行(KB/s)": round(88.5 + second / 2, 2),
        "进程数": 1,
        "上下文切换/秒": 1500 + second * 10,
    }
    for second in range(6)
]


def test_fixed_headers_cover_core_and_io_fields():
    # 表头必须能被页面 io.normalize_scene 直接识别
    for header in ("时间戳", "CPU%", "内存占用（MB）", "系统总内存", "内存占用百分比", "线程", "句柄", "io"):
        assert header in FIXED_HEADERS
    for header in ("磁盘读(MB/s)", "磁盘写(MB/s)", "IOPS", "网络上行(KB/s)", "网络下行(KB/s)"):
        assert header in FIXED_HEADERS


def test_output_columns_are_unique_and_ordered():
    columns = output_columns()
    assert len(columns) == len(set(columns))
    assert columns[: len(FIXED_HEADERS)] == FIXED_HEADERS
    assert columns[len(FIXED_HEADERS):] == BUSINESS_HEADERS


def _fake_psutil(processes: list | None = None, **overrides):
    """构造一个最小可用的假 psutil 模块，返回 (模块, 还原函数)。"""
    import sys
    import types

    fake = types.ModuleType("psutil")
    fake.cpu_percent = lambda interval=None: 12.5
    fake.virtual_memory = lambda: types.SimpleNamespace(total=24 * 1024**3, used=8 * 1024**3)
    fake.disk_io_counters = lambda: types.SimpleNamespace(
        read_bytes=0, write_bytes=0, read_count=0, write_count=0
    )
    fake.net_io_counters = lambda: types.SimpleNamespace(bytes_sent=0, bytes_recv=0)
    fake.cpu_stats = lambda: types.SimpleNamespace(ctx_switches=1000)
    fake.process_iter = lambda attrs=None: list(processes or [])
    for key, value in overrides.items():
        setattr(fake, key, value)

    original = sys.modules.get("psutil")
    sys.modules["psutil"] = fake

    def restore():
        if original is None:
            sys.modules.pop("psutil", None)
        else:  # pragma: no cover
            sys.modules["psutil"] = original

    return fake, restore


def test_system_mode_fills_threads_and_handles():
    """回归：整机口径下线程/句柄不能再是空列（二者都是核心指标）。"""
    from perf_analyzer import collector

    class FakeProc:
        def __init__(self, threads, fds):
            self.info = {"num_threads": threads}
            self._fds = fds

        def num_fds(self):
            return self._fds

    fake, restore = _fake_psutil(processes=[FakeProc(10, 100), FakeProc(5, 50)])
    try:
        row = collector.Sampler(fake, [], interval=1.0).sample()
    finally:
        restore()

    assert row["线程"] == 15
    assert row["进程数"] == 1


def test_system_mode_handles_from_file_nr(tmp_path, monkeypatch):
    """整机句柄优先读 /proc/sys/fs/file-nr（系统级已分配文件句柄）。"""
    from perf_analyzer import collector

    file_nr = tmp_path / "file-nr"
    file_nr.write_text("12345\t0\t0\n", encoding="utf-8")

    class FakeProc:
        info = {"num_threads": 7}

        def num_fds(self):
            raise PermissionError("denied")

    fake, restore = _fake_psutil(processes=[FakeProc()])
    original_read = collector.Path.read_text

    def fake_read_text(self, *args, **kwargs):
        if str(self) == "/proc/sys/fs/file-nr":
            return file_nr.read_text(encoding="utf-8")
        return original_read(self, *args, **kwargs)

    monkeypatch.setattr(collector.Path, "read_text", fake_read_text)
    try:
        sampler = collector.Sampler(fake, [], interval=1.0)
        handles = sampler._system_handles()
        row = sampler.sample()
    finally:
        restore()

    assert handles == 12345
    assert row["句柄"] == 12345


def test_system_mode_memory_uses_system_used():
    """整机口径下内存占用取系统已用量，内存占比随之可算。"""
    from perf_analyzer import collector

    fake, restore = _fake_psutil(virtual_memory=lambda: types.SimpleNamespace(total=16 * 1024**3, used=4 * 1024**3))
    try:
        row = collector.Sampler(fake, [], interval=1.0).sample()
    finally:
        restore()

    assert row["内存占用（MB）"] == 4096.0
    assert row["系统总内存"] == 16384.0
    assert abs(row["内存占用百分比"] - 25.0) < 1e-6


def test_process_mode_handle_denial_does_not_truncate_threads():
    """进程树里某个进程读不到 fd，不该让整棵树的线程数也被截断。"""
    from perf_analyzer import collector

    class Denied:
        pid = 1
        name = "root-proc"

        def is_running(self):
            return True

        def memory_full_info(self):
            return types.SimpleNamespace(rss=100 * 1024 * 1024, uss=50 * 1024 * 1024)

        def num_threads(self):
            return 20

        def num_fds(self):
            raise PermissionError("denied")

    class Ok:
        pid = 2
        name = "app"

        def is_running(self):
            return True

        def memory_full_info(self):
            return types.SimpleNamespace(rss=200 * 1024 * 1024, uss=100 * 1024 * 1024)

        def num_threads(self):
            return 5

        def num_fds(self):
            return 42

    fake, restore = _fake_psutil(processes=[Denied(), Ok()])
    try:
        row = collector.Sampler(fake, [Denied(), Ok()], interval=1.0).sample()
    finally:
        restore()

    assert row["线程"] == 25
    assert row["句柄"] is None  # 有进程读不到 fd，宁缺勿错
    assert row["进程数"] == 2


def test_describe_total_memory_uses_gb_unit():
    """回归：系统内存展示值必须是 GB，而不是把 MB 数当 GB（曾出现 25292 GB）。"""
    from perf_analyzer import collector

    _, restore = _fake_psutil()
    try:
        value = collector.describe_total_memory()
    finally:
        restore()
    assert value == "24.0 GB"


def test_build_system_info_rows():
    info = build_system_info("DemoApp", "1.0.0", "Lab PC", 5.0, "java（进程数 2）", extra={"备注": "压测"})
    assert list(info.columns) == ["项目", "值"]
    assert "应用名称" in set(info["项目"])
    row = dict(zip(info["项目"], info["值"]))
    assert row["采样间隔"] == "5 秒"
    assert row["采集目标"] == "java（进程数 2）"
    assert row["备注"] == "压测"


def test_map_metrics_snapshot_english_and_chinese_keys():
    mapped = map_metrics_snapshot(
        {
            "qps": 120.5,
            "latency_ms": 235.8,
            "concurrency": 80,
            "requests": 1000,
            "failed_requests": 3,
            "retry_count": 1,
            "heap_mb": 512,
            "unknown_field": "忽略我",
        }
    )
    assert mapped["QPS"] == 120.5
    assert mapped["响应时间(ms)"] == 235.8
    assert mapped["并发数"] == 80
    assert mapped["请求数"] == 1000
    # 关键回归：failed_requests 不能被 "requests" 子串抢走
    assert mapped["失败请求数"] == 3
    assert mapped["重试次数"] == 1
    assert mapped["堆内存(MB)"] == 512
    assert "unknown_field" not in mapped.values()


def test_map_metrics_snapshot_percent_not_memory_mb():
    mapped = map_metrics_snapshot({"memory_percent": 12.5})
    assert mapped["内存占用百分比"] == 12.5
    assert "内存占用（MB）" not in mapped


def test_read_metrics_snapshot_json_and_jsonl(tmp_path):
    obj = tmp_path / "metrics.json"
    obj.write_text(json.dumps({"qps": 10, "并发数": 5}), encoding="utf-8")
    assert read_metrics_snapshot(obj) == {"qps": 10, "并发数": 5}

    lines = tmp_path / "metrics.jsonl"
    lines.write_text(
        "\n".join([json.dumps({"qps": 1}), json.dumps({"qps": 2}), "{}"]),
        encoding="utf-8",
    )
    assert read_metrics_snapshot(lines) == {}

    lines.write_text(
        "\n".join([json.dumps({"qps": 1}), json.dumps({"qps": 99})]),
        encoding="utf-8",
    )
    assert read_metrics_snapshot(lines)["qps"] == 99

    assert read_metrics_snapshot(tmp_path / "missing.json") == {}
    broken = tmp_path / "broken.json"
    broken.write_text("{oops", encoding="utf-8")
    assert read_metrics_snapshot(broken) == {}


def test_new_sheet_name_avoids_collision():
    assert new_sheet_name([], "采集场景1") == "采集场景1"
    assert new_sheet_name(["采集场景1"], "采集场景1") == "采集场景1-2"
    assert new_sheet_name(["采集场景1", "采集场景1-2"], "采集场景1") == "采集场景1-3"


def test_build_frame_column_order_and_nan():
    frame = build_frame([{"时间戳": "2026-06-13 10:00:00", "CPU%": 1.0}], ["QPS"])
    assert list(frame.columns) == [*FIXED_HEADERS, "QPS"]
    assert frame["CPU%"].iloc[0] == 1.0
    assert frame["QPS"].isna().all()


def test_write_workbook_creates_then_appends(tmp_path):
    path = tmp_path / "collected.xlsx"
    info = build_system_info("App", "1.0", "PC", 5.0, "整机")
    frame = build_frame(SAMPLE_ROWS)

    write_workbook(path, info, "轻载", frame)
    assert [*pd_sheetnames(path)] == ["系统信息", "轻载"]

    # 追加第二个场景，第一个不动
    write_workbook(path, info, "重载", frame)
    assert [*pd_sheetnames(path)] == ["系统信息", "轻载", "重载"]

    # 同名场景默认换新 sheet，不覆盖历史
    _, target = write_workbook(path, info, "重载", frame)
    assert target == "重载-2"
    # 显式覆盖则回到原名
    _, target = write_workbook(path, info, "重载", frame, overwrite_sheet=True)
    assert target == "重载"


def test_collected_workbook_is_readable_by_the_tool(tmp_path):
    """端到端：采集器写出的 Excel 必须被页面自身的数据装载逻辑识别。"""
    path = tmp_path / "collected.xlsx"
    write_workbook(
        path,
        build_system_info("DemoApp", "1.0.0", "Lab PC", 5.0, "整机"),
        "采集场景1",
        build_frame(SAMPLE_ROWS, ["QPS", "响应时间(ms)"]),
    )

    system_info, data = read_single_workbook("采集数据", None, (), PERFORMANCE_MODE, path)

    assert set(data["scene"]) == {"采集场景1"}
    assert set(data["source"]) == {"采集数据"}
    assert len(data) == len(SAMPLE_ROWS)
    assert data["timestamp"].notna().all()

    for column in (
        "cpu_percent",
        "memory_mb",
        "total_memory_mb",
        "memory_percent",
        "threads",
        "handles",
        "io",
        "disk_read_mbps",
        "disk_write_mbps",
        "iops",
        "network_up_kbps",
        "network_down_kbps",
    ):
        assert column in data.columns, column
        assert data[column].notna().all(), column

    # 业务列来自 --metrics-file 挂接
    assert data["qps"].notna().any() or "qps" in data.columns
    assert "上下文切换/秒" in {label for label in data.columns} or any(
        column.startswith("extra::") for column in data.columns
    )
    assert not system_info.empty
    assert "数据源" in system_info.columns


def pd_sheetnames(path: Path):
    from openpyxl import load_workbook

    return load_workbook(path, read_only=True).sheetnames


def _collected_frames(scenes: int = 2, points: int = 60) -> list[dict]:
    """构造两个场景的采集数据：资源列齐全 + 业务列来自 metrics-file。"""
    frames = []
    for scene_index in range(scenes):
        load = 30 + scene_index * 25
        rows = []
        for index in range(points):
            cpu = load + 6 * ((index % 7) - 3)
            memory = 700 + scene_index * 200 + index * 2.5  # 有明显增长趋势
            rows.append(
                {
                    "时间戳": f"2026-06-13 1{scene_index}:{index // 6:02d}:{index % 6 * 10:02d}",
                    "CPU%": round(cpu, 2),
                    "内存占用（MB）": round(memory, 2),
                    "系统总内存": 16384,
                    "内存占用百分比": round(memory / 16384 * 100, 4),
                    "线程": 40 + scene_index * 10,
                    "句柄": 800 + index,
                    "io": round(1.0 + index / 20, 3),
                    "磁盘读(MB/s)": round(0.4 + index / 40, 3),
                    "磁盘写(MB/s)": round(0.2 + index / 60, 3),
                    "IOPS": 40 + index,
                    "网络上行(KB/s)": 12.5,
                    "网络下行(KB/s)": 88.5,
                    "进程数": 1,
                    "上下文切换/秒": 1500 + index * 10,
                    "QPS": 120 - scene_index * 20,
                    "并发数": 80 - scene_index * 10,
                    "响应时间(ms)": round(180 + cpu, 1),
                    "请求数": 1200,
                    "失败请求数": index % 5,
                    "重试次数": index % 3,
                }
            )
        frames.append((f"采集场景{scene_index + 1}", build_frame(rows, ["QPS", "并发数", "响应时间(ms)", "请求数", "失败请求数", "重试次数"])))
    return frames


def test_collected_data_flows_through_real_analysis(tmp_path):
    """采集器产出的数据必须能直接喂进页面自己的分析链路，而不只是列名对得上。"""
    from perf_analyzer.analysis import (
        alert_continuity_summary,
        configuration_recommendation,
        leak_analysis,
        resource_efficiency_analysis,
        summarize,
    )
    from perf_analyzer.constants import PERFORMANCE_MODE
    from perf_analyzer.io import read_single_workbook
    from perf_analyzer.metrics import discover_metrics

    path = tmp_path / "collected.xlsx"
    info = build_system_info("DemoApp", "1.0.0", "Lab PC", 5.0, "MyApp（进程数 1）")
    for name, frame in _collected_frames():
        write_workbook(path, info, name, frame)

    _, data = read_single_workbook("采集冒烟", None, (), PERFORMANCE_MODE, path)
    assert len(data) == 120
    assert set(data["scene"]) == {"采集场景1", "采集场景2"}

    metrics = discover_metrics(data)
    keys = {metric.key for metric in metrics}
    # 采集器写的固定列 + 业务列都应该被识别成指标
    for key in ("cpu_percent", "memory_mb", "threads", "handles", "io", "qps", "concurrency", "response_time_ms"):
        assert key in keys, key
    thresholds = {metric.key: metric.default_threshold for metric in metrics}

    summary = summarize(data, thresholds, metrics)
    assert not summary.empty and len(summary) == 2
    # 列名按指标 label 生成：CPU→CPU均值/CPUP95，内存→内存占用均值/内存占用峰值
    assert "CPU均值" in summary.columns
    assert "CPUP95" in summary.columns
    assert "内存占用峰值" in summary.columns
    assert summary["样本数"].sum() == 120

    leaks = leak_analysis(data, 2.0)
    assert set(leaks["场景"]) == {"采集场景1", "采集场景2"}
    assert set(leaks["判断"]) <= {"疑似泄漏", "轻微增长", "无明显增长", "数据不足"}
    assert (leaks["内存斜率(MB/分钟)"] > 0).all()

    alerts = alert_continuity_summary(data, thresholds, metrics)
    assert not alerts.empty
    assert {"指标", "超阈次数", "最大连续超阈"} <= set(alerts.columns)

    recommendation, notes = configuration_recommendation(
        data, test_cpu_cores=8, target_cpu_utilization=70.0, minimum_cpu_cores=2,
        system_reserved_memory_gb=2.0, minimum_memory_gb=4, minimum_free_disk_gb=10,
        recommendation_factor=1.5, slope_threshold=2.0,
    )
    assert list(recommendation["配置档位"]) == ["最低可运行", "推荐配置"]
    assert recommendation["CPU核数"].iloc[0] >= 2
    assert notes and any("CPU估算" in note for note in notes)

    efficiency = resource_efficiency_analysis(data)
    assert not efficiency.empty and len(efficiency) == 2


@pytest.mark.parametrize(
    "raw,expected",
    [
        ({"请求数": 10}, "请求数"),
        ({"request_count": 10}, "请求数"),
        ({"请求吞吐": 10}, "QPS"),
    ],
)
def test_header_alias_examples(raw, expected):
    assert expected in map_metrics_snapshot(raw)
