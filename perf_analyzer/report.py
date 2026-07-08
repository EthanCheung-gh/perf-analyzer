from __future__ import annotations

import html
import io
import os
import zipfile
from datetime import datetime

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

from .alerts import alert_segments, format_value, metric_alert_mask
from .analysis import (
    alert_continuity_summary,
    bottleneck_hints,
    capacity_estimation,
    leak_analysis,
    radar_profile_data,
    release_gate_analysis,
    resource_efficiency_analysis,
    risk_matrix_data,
    scene_decision_analysis,
    stability_score_analysis,
    summarize,
    trend_analysis,
)
from .models import Metric
from .ui_helpers import with_display_label


def _format_report_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def _count_alerts(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> int:
    total = 0
    for metric in metrics:
        if metric.key in data:
            total += int(metric_alert_mask(data[metric.key], metric, thresholds, threshold_bounds).sum())
    return total


def build_conclusions(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    slope_threshold: float,
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> list[str]:
    conclusions = []
    if data.empty:
        return ["未读取到可分析的场景数据。"]
    alerts = _count_alerts(data, thresholds, metrics, threshold_bounds)
    if alerts:
        conclusions.append(f"当前阈值下共有 {alerts} 个超阈采样点，建议优先查看“异常”页定位时间和场景。")
        alert_summary = alert_continuity_summary(data, thresholds, metrics, threshold_bounds)
        if not alert_summary.empty:
            top_ratio = alert_summary.iloc[0]
            top_duration = alert_summary.sort_values("最大连续超阈(分钟)", ascending=False).iloc[0]
            conclusions.append(
                f"超阈占比最高的是 {top_ratio['数据源']} / {top_ratio['场景']} 的 {top_ratio['指标']}，"
                f"占比 {top_ratio['超阈占比(%)']:.2f}%。"
            )
            conclusions.append(
                f"最长连续超阈出现在 {top_duration['数据源']} / {top_duration['场景']} 的 {top_duration['指标']}，"
                f"持续 {top_duration['最大连续超阈']}。"
            )
    else:
        conclusions.append("当前阈值下未发现超阈采样点。")

    for metric in metrics:
        if data[metric.key].dropna().empty:
            continue
        idx = data[metric.key].idxmax()
        row = data.loc[idx]
        conclusions.append(
            f"{metric.label}峰值出现在 {row['source']} / {row['scene']}，峰值为 {format_value(row[metric.key], metric.unit)}。"
        )

    leaks = leak_analysis(data, slope_threshold)
    risky = leaks[leaks["判断"] == "疑似泄漏"]
    if not risky.empty:
        names = "、".join(f"{row['数据源']} / {row['场景']}" for _, row in risky.iterrows())
        conclusions.append(f"内存趋势检测发现疑似泄漏场景：{names}。")
    else:
        conclusions.append("内存趋势检测未发现明显泄漏场景。")
    return conclusions


def build_generic_conclusions(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> list[str]:
    conclusions = []
    if data.empty:
        return ["未读取到可分析的时序数据。"]
    alerts = _count_alerts(data, thresholds, metrics, threshold_bounds)
    if alerts:
        conclusions.append(f"当前阈值下共有 {alerts} 个超阈采样点，建议查看“异常”页定位时间和分组。")
        alert_summary = alert_continuity_summary(data, thresholds, metrics, threshold_bounds)
        if not alert_summary.empty:
            top_ratio = alert_summary.iloc[0]
            conclusions.append(
                f"超阈占比最高的是 {top_ratio['数据源']} / {top_ratio['场景']} 的 {top_ratio['指标']}，"
                f"占比 {top_ratio['超阈占比(%)']:.2f}%。"
            )
    else:
        conclusions.append("当前阈值下未发现超阈采样点。")

    for metric in metrics:
        if metric.key not in data or data[metric.key].dropna().empty:
            continue
        idx = data[metric.key].idxmax()
        row = data.loc[idx]
        conclusions.append(
            f"{metric.label}峰值出现在 {row['source']} / {row['scene']}，峰值为 {format_value(row[metric.key], metric.unit)}。"
        )
    return conclusions


def dataframe_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    display = df.copy()
    for column in display.columns:
        display[column] = display[column].map(_format_report_value)
    header = "| " + " | ".join(display.columns) + " |"
    separator = "| " + " | ".join("---" for _ in display.columns) + " |"
    rows = ["| " + " | ".join(row) + " |" for row in display.astype(str).to_numpy()]
    return "\n".join([header, separator, *rows])


def _top_rows(df: pd.DataFrame, sort_column: str | None = None, limit: int = 10, ascending: bool = False) -> pd.DataFrame:
    if df.empty:
        return df
    result = df.copy()
    if sort_column and sort_column in result:
        result = result.sort_values(sort_column, ascending=ascending)
    return result.head(limit)


def _select_existing_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    if df.empty:
        return df
    selected = [column for column in columns if column in df.columns]
    return df[selected] if selected else df


def _safe_sheet_name(name: str, used: set[str]) -> str:
    illegal_chars = set("[]:*?/\\")
    cleaned = "".join("_" if char in illegal_chars else char for char in name).strip() or "Sheet"
    base = cleaned[:31]
    candidate = base
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate)
    return candidate


def _matplotlib_png(draw: object, width: float = 11.0, height: float = 5.8) -> bytes:
    try:
        os.environ.setdefault("MPLCONFIGDIR", "/tmp/perf-analyzer-matplotlib")
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ModuleNotFoundError as exc:
        raise RuntimeError("未安装 matplotlib，无法生成 PDF 静态图表。请先执行 pip install matplotlib。") from exc

    plt.rcParams["font.sans-serif"] = [
        "Noto Sans CJK SC",
        "Microsoft YaHei",
        "SimHei",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False
    fig = plt.figure(figsize=(width, height), dpi=160)
    try:
        draw(fig, plt)
        output = io.BytesIO()
        fig.tight_layout()
        fig.savefig(output, format="png", bbox_inches="tight", facecolor="white")
        return output.getvalue()
    finally:
        plt.close(fig)


def _append_matplotlib_image(images: list[tuple[str, bytes]], title: str, builder: object, errors: list[str]) -> None:
    try:
        image = builder()
        if image:
            images.append((title, image))
    except RuntimeError as exc:
        errors.append(str(exc))


def _build_risk_matrix_image(decisions: pd.DataFrame) -> bytes | None:
    matrix = risk_matrix_data(decisions)
    if matrix.empty:
        return None
    matrix_pivot = matrix.pivot(index="场景标识", columns="风险项", values="风险分").fillna(0)

    def draw(fig: object, plt: object) -> None:
        ax = fig.add_subplot(111)
        data_values = matrix_pivot.values.astype(float)
        image = ax.imshow(data_values, vmin=0, vmax=2, cmap="RdYlGn_r", aspect="auto")
        ax.set_title("场景风险矩阵")
        ax.set_xticks(range(len(matrix_pivot.columns)))
        ax.set_xticklabels(matrix_pivot.columns, rotation=30, ha="right")
        ax.set_yticks(range(len(matrix_pivot.index)))
        ax.set_yticklabels(matrix_pivot.index)
        for row_index in range(data_values.shape[0]):
            for col_index in range(data_values.shape[1]):
                ax.text(col_index, row_index, f"{data_values[row_index, col_index]:.0f}", ha="center", va="center", fontsize=8)
        fig.colorbar(image, ax=ax, label="风险分")

    return _matplotlib_png(draw)


def _build_radar_image(data: pd.DataFrame, decisions: pd.DataFrame, thresholds: dict[str, float]) -> bytes | None:
    radar_data = radar_profile_data(data, decisions, thresholds)
    if radar_data.empty:
        return None
    scene_options = list(radar_data["场景标识"].drop_duplicates())[:3]
    dimensions = list(radar_data[radar_data["场景标识"] == scene_options[0]]["维度"])
    if not dimensions:
        return None
    angles = np.linspace(0, 2 * np.pi, len(dimensions), endpoint=False).tolist()
    angles.append(angles[0])

    def draw(fig: object, plt: object) -> None:
        ax = fig.add_subplot(111, polar=True)
        for scene_id in scene_options:
            scene_profile = radar_data[radar_data["场景标识"] == scene_id]
            values = list(scene_profile["风险分"].astype(float))
            if not values:
                continue
            values.append(values[0])
            ax.plot(angles, values, linewidth=1.8, label=scene_id)
            ax.fill(angles, values, alpha=0.12)
        ax.set_title("场景风险画像")
        ax.set_ylim(0, 100)
        ax.set_xticks(angles[:-1])
        ax.set_xticklabels(dimensions)
        ax.legend(loc="upper right", bbox_to_anchor=(1.28, 1.1), fontsize=8)

    return _matplotlib_png(draw)


def _build_alert_timeline_image(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> bytes | None:
    alert_frames = []
    for metric in metrics:
        if metric.key not in data:
            continue
        mask = metric_alert_mask(data[metric.key], metric, thresholds, threshold_bounds)
        alert = data[mask][["timestamp", "source", "scene", metric.key]].copy()
        if alert.empty:
            continue
        alert["指标"] = metric.label
        alert_frames.append(alert[["timestamp", "source", "scene", "指标"]])
    if not alert_frames:
        return None
    alerts = with_display_label(pd.concat(alert_frames, ignore_index=True).sort_values("timestamp"))
    scenes = list(alerts["display_scene"].drop_duplicates())
    metrics_labels = list(alerts["指标"].drop_duplicates())
    scene_to_y = {scene: index for index, scene in enumerate(scenes)}
    color_map = {label: index for index, label in enumerate(metrics_labels)}

    def draw(fig: object, plt: object) -> None:
        ax = fig.add_subplot(111)
        colors = plt.get_cmap("tab10")
        for label in metrics_labels:
            subset = alerts[alerts["指标"] == label]
            ax.scatter(
                subset["timestamp"],
                subset["display_scene"].map(scene_to_y),
                s=36,
                color=colors(color_map[label] % 10),
                label=label,
                alpha=0.82,
            )
        ax.set_title("异常时间轴")
        ax.set_xlabel("时间")
        ax.set_yticks(range(len(scenes)))
        ax.set_yticklabels(scenes)
        ax.legend(loc="best", fontsize=8)
        fig.autofmt_xdate(rotation=20)

    return _matplotlib_png(draw)


def _build_capacity_curve_image(data: pd.DataFrame) -> bytes | None:
    capacity = capacity_estimation(data, 8, 70.0, [50, 100, 200], 2.0)
    if capacity.empty:
        return None

    def draw(fig: object, plt: object) -> None:
        ax = fig.add_subplot(111)
        ax.plot(capacity["目标吞吐"], capacity["建议CPU核数"], marker="o", color="#1f77b4", label="建议CPU核数")
        ax.set_title("容量曲线")
        ax.set_xlabel("目标吞吐")
        ax.set_ylabel("建议CPU核数")
        if "内存需求(GB)" in capacity and capacity["内存需求(GB)"].notna().any():
            ax2 = ax.twinx()
            ax2.plot(capacity["目标吞吐"], capacity["内存需求(GB)"], marker="o", color="#2ca02c", label="内存需求(GB)")
            ax2.set_ylabel("内存需求(GB)")
        ax.grid(True, alpha=0.25)

    return _matplotlib_png(draw)


def _build_top_trend_image(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> bytes | None:
    if data.empty:
        return None
    selected_metric: Metric | None = None
    best_alerts = -1
    for metric in metrics:
        if metric.key not in data or data[metric.key].dropna().empty:
            continue
        count = int(metric_alert_mask(data[metric.key], metric, thresholds, threshold_bounds).sum())
        if count > best_alerts:
            best_alerts = count
            selected_metric = metric
    if selected_metric is None:
        return None
    filtered = with_display_label(data.copy()).sort_values("timestamp")

    def draw(fig: object, plt: object) -> None:
        ax = fig.add_subplot(111)
        for display_scene, group in filtered.groupby("display_scene", sort=False):
            ax.plot(group["timestamp"], group[selected_metric.key], linewidth=1.8, label=display_scene)
        if selected_metric.key in thresholds:
            ax.axhline(thresholds[selected_metric.key], color="#d62728", linestyle="--", linewidth=1.2, label="阈值")
        ax.set_title(f"关键指标趋势：{selected_metric.label}")
        ax.set_xlabel("时间")
        ax.set_ylabel(f"{selected_metric.label}({selected_metric.unit})")
        ax.legend(loc="best", fontsize=8)
        ax.grid(True, alpha=0.25)
        fig.autofmt_xdate(rotation=20)

    return _matplotlib_png(draw)


def _build_risk_matrix_figure(decisions: pd.DataFrame) -> go.Figure | None:
    matrix = risk_matrix_data(decisions)
    if matrix.empty:
        return None
    matrix_pivot = matrix.pivot(index="场景标识", columns="风险项", values="风险分").fillna(0)
    fig = go.Figure(
        data=go.Heatmap(
            z=matrix_pivot.values,
            x=matrix_pivot.columns,
            y=matrix_pivot.index,
            zmin=0,
            zmax=2,
            colorscale=[[0, "#2ca02c"], [0.5, "#ffbf00"], [1, "#d62728"]],
            colorbar=dict(title="风险分", tickvals=[0, 1, 2], ticktext=["低", "中", "高"]),
            text=matrix_pivot.values,
            texttemplate="%{text:.0f}",
        )
    )
    fig.update_layout(title="场景风险矩阵", height=420, margin=dict(l=30, r=30, t=70, b=40))
    return fig


def _build_radar_figure(data: pd.DataFrame, decisions: pd.DataFrame, thresholds: dict[str, float]) -> go.Figure | None:
    radar_data = radar_profile_data(data, decisions, thresholds)
    if radar_data.empty:
        return None
    scene_options = list(radar_data["场景标识"].drop_duplicates())[:3]
    fig = go.Figure()
    for scene_id in scene_options:
        scene_profile = radar_data[radar_data["场景标识"] == scene_id]
        dimensions = list(scene_profile["维度"])
        values = list(scene_profile["风险分"])
        if dimensions:
            fig.add_trace(
                go.Scatterpolar(
                    r=[*values, values[0]],
                    theta=[*dimensions, dimensions[0]],
                    fill="toself",
                    name=scene_id,
                )
            )
    if not fig.data:
        return None
    fig.update_layout(
        title="场景风险画像",
        polar=dict(radialaxis=dict(visible=True, range=[0, 100])),
        height=520,
        margin=dict(l=30, r=30, t=70, b=40),
    )
    return fig


def _build_alert_timeline_figure(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> go.Figure | None:
    alert_frames = []
    for metric in metrics:
        if metric.key not in data:
            continue
        mask = metric_alert_mask(data[metric.key], metric, thresholds, threshold_bounds)
        alert = data[mask][["timestamp", "source", "scene", metric.key]].copy()
        if alert.empty:
            continue
        alert["指标"] = metric.label
        alert["值"] = alert[metric.key]
        alert_frames.append(alert[["timestamp", "source", "scene", "指标", "值"]])
    if not alert_frames:
        return None
    alerts = pd.concat(alert_frames, ignore_index=True).sort_values("timestamp")
    timeline = with_display_label(alerts)
    fig = px.scatter(
        timeline,
        x="timestamp",
        y="display_scene",
        color="指标",
        hover_data=["source", "scene", "值"],
        labels={"timestamp": "时间", "display_scene": "场景"},
        title="异常时间轴",
    )
    fig.update_traces(marker=dict(size=9))
    fig.update_layout(height=460, margin=dict(l=30, r=30, t=70, b=40))
    return fig


def _build_capacity_curve_figure(data: pd.DataFrame) -> go.Figure | None:
    capacity = capacity_estimation(data, 8, 70.0, [50, 100, 200], 2.0)
    if capacity.empty:
        return None
    fig = go.Figure()
    fig.add_trace(
        go.Scatter(
            x=capacity["目标吞吐"],
            y=capacity["建议CPU核数"],
            mode="lines+markers",
            name="建议CPU核数",
            line=dict(color="#1f77b4"),
        )
    )
    if "内存需求(GB)" in capacity and capacity["内存需求(GB)"].notna().any():
        fig.add_trace(
            go.Scatter(
                x=capacity["目标吞吐"],
                y=capacity["内存需求(GB)"],
                mode="lines+markers",
                name="内存需求(GB)",
                yaxis="y2",
                line=dict(color="#2ca02c"),
            )
        )
    fig.update_layout(
        title="容量曲线",
        xaxis=dict(title="目标吞吐"),
        yaxis=dict(title="建议CPU核数"),
        yaxis2=dict(title="内存需求(GB)", overlaying="y", side="right"),
        height=460,
        margin=dict(l=30, r=30, t=70, b=40),
    )
    return fig


def _build_top_trend_figure(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> go.Figure | None:
    if data.empty:
        return None
    selected_metric: Metric | None = None
    best_alerts = -1
    for metric in metrics:
        if metric.key not in data or data[metric.key].dropna().empty:
            continue
        count = int(metric_alert_mask(data[metric.key], metric, thresholds, threshold_bounds).sum())
        if count > best_alerts:
            best_alerts = count
            selected_metric = metric
    if selected_metric is None:
        return None
    filtered = with_display_label(data.copy())
    fig = px.line(
        filtered.sort_values("timestamp"),
        x="timestamp",
        y=selected_metric.key,
        color="display_scene",
        labels={"timestamp": "时间", selected_metric.key: f"{selected_metric.label}({selected_metric.unit})", "display_scene": "场景"},
        title=f"关键指标趋势：{selected_metric.label}",
    )
    if selected_metric.key in thresholds:
        fig.add_hline(y=thresholds[selected_metric.key], line_dash="dash", line_color="#d62728", annotation_text="阈值")
    fig.update_layout(height=460, margin=dict(l=30, r=30, t=70, b=40))
    return fig


def build_report_figure_images(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    slope_threshold: float,
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> tuple[list[tuple[str, bytes]], list[str]]:
    decisions = scene_decision_analysis(data, thresholds, metrics, slope_threshold, threshold_bounds)
    images: list[tuple[str, bytes]] = []
    errors: list[str] = []
    _append_matplotlib_image(images, "场景风险矩阵", lambda: _build_risk_matrix_image(decisions), errors)
    _append_matplotlib_image(images, "场景风险画像", lambda: _build_radar_image(data, decisions, thresholds), errors)
    _append_matplotlib_image(images, "异常时间轴", lambda: _build_alert_timeline_image(data, thresholds, metrics, threshold_bounds), errors)
    _append_matplotlib_image(images, "容量曲线", lambda: _build_capacity_curve_image(data), errors)
    _append_matplotlib_image(images, "关键指标趋势", lambda: _build_top_trend_image(data, thresholds, metrics, threshold_bounds), errors)
    return images, errors


def build_generic_report_figure_images(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> tuple[list[tuple[str, bytes]], list[str]]:
    images: list[tuple[str, bytes]] = []
    errors: list[str] = []
    _append_matplotlib_image(images, "异常时间轴", lambda: _build_alert_timeline_image(data, thresholds, metrics, threshold_bounds), errors)
    _append_matplotlib_image(images, "关键指标趋势", lambda: _build_top_trend_image(data, thresholds, metrics, threshold_bounds), errors)
    return images, errors


def _markdown_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_markdown_table_separator(line: str) -> bool:
    cells = _markdown_table_row(line)
    return bool(cells) and all(cell.replace("-", "").replace(":", "").strip() == "" and "-" in cell for cell in cells)


def _markdown_blocks(markdown_text: str) -> list[dict[str, object]]:
    lines = markdown_text.splitlines()
    blocks: list[dict[str, object]] = []
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            blocks.append({"type": "heading", "level": 1, "text": line[2:].strip()})
            index += 1
            continue
        if line.startswith("## "):
            blocks.append({"type": "heading", "level": 2, "text": line[3:].strip()})
            index += 1
            continue
        if line.startswith("- "):
            items = []
            while index < len(lines) and lines[index].startswith("- "):
                items.append(lines[index][2:].strip())
                index += 1
            blocks.append({"type": "list", "items": items})
            continue
        if line.startswith("|") and index + 1 < len(lines) and _is_markdown_table_separator(lines[index + 1]):
            headers = _markdown_table_row(line)
            rows = []
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(_markdown_table_row(lines[index]))
                index += 1
            blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue
        paragraphs = [line.strip()]
        index += 1
        while index < len(lines):
            next_line = lines[index].rstrip()
            if not next_line.strip() or next_line.startswith("# ") or next_line.startswith("## ") or next_line.startswith("- ") or next_line.startswith("|"):
                break
            paragraphs.append(next_line.strip())
            index += 1
        blocks.append({"type": "paragraph", "text": " ".join(paragraphs)})
    return blocks


def build_report_markdown(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    slope_threshold: float,
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
    config_table: pd.DataFrame | None = None,
    config_notes: list[str] | None = None,
) -> str:
    summary = summarize(data, thresholds, metrics, threshold_bounds)
    leaks = leak_analysis(data, slope_threshold)
    trends = trend_analysis(data, metrics)
    alert_summary = alert_continuity_summary(data, thresholds, metrics, threshold_bounds)
    segments = alert_segments(data, thresholds, metrics, threshold_bounds)
    efficiency = resource_efficiency_analysis(data)
    decisions = scene_decision_analysis(data, thresholds, metrics, slope_threshold, threshold_bounds)
    gate_status, gate_rules = release_gate_analysis(
        decisions,
        max_allowed_level="中风险",
        fail_on_leak=True,
        fail_on_any_alert=False,
        fail_on_errors=False,
    )
    stability = stability_score_analysis(decisions)
    bottlenecks = bottleneck_hints(data, thresholds, metrics)
    conclusions = build_conclusions(data, thresholds, slope_threshold, metrics, threshold_bounds)
    lines = [
        "# 应用资源消耗分析报告",
        "",
        "## 测试结论",
        "",
        *[f"- {item}" for item in conclusions],
        f"- 发布门禁默认判断：{gate_status}。",
        "",
        "## 决策分析",
        "",
        dataframe_to_markdown(decisions) if not decisions.empty else "无场景风险评级数据。",
        "",
        "## 发布门禁",
        "",
        dataframe_to_markdown(gate_rules) if not gate_rules.empty else "无发布门禁数据。",
        "",
        "## 稳定性评分",
        "",
        dataframe_to_markdown(stability) if not stability.empty else "无稳定性评分数据。",
        "",
        "## 瓶颈提示",
        "",
        dataframe_to_markdown(bottlenecks) if not bottlenecks.empty else "无瓶颈提示数据。",
        "",
        "## 汇总数据",
        "",
        dataframe_to_markdown(summary) if not summary.empty else "无汇总数据。",
        "",
        "## 告警持续性",
        "",
        dataframe_to_markdown(alert_summary.drop(columns=["最大连续超阈(分钟)"], errors="ignore")) if not alert_summary.empty else "当前阈值下无超阈持续性数据。",
        "",
        "## 异常段",
        "",
        dataframe_to_markdown(segments.drop(columns=["指标key", "持续时长(分钟)"], errors="ignore")) if not segments.empty else "当前阈值下无异常段数据。",
        "",
        "## 内存趋势检测",
        "",
        dataframe_to_markdown(leaks) if not leaks.empty else "无内存趋势数据。",
        "",
        "## 全指标趋势斜率",
        "",
        dataframe_to_markdown(trends) if not trends.empty else "无趋势斜率数据。",
        "",
        "## 资源效率",
        "",
        dataframe_to_markdown(efficiency) if not efficiency.empty else "无资源效率数据。",
        "",
        "## 配置建议",
        "",
        dataframe_to_markdown(config_table) if config_table is not None and not config_table.empty else "未生成配置建议。",
        "",
        *[f"- {item}" for item in config_notes or []],
        "",
        "## 分析口径",
        "",
        "- 均值：场景内全部采样点平均值",
        "- 峰值：场景内最大采样值",
        "- P90 / P95 / P99：对应分位值，用于观察高水位和尾部风险",
        "- 超阈次数：按单阈值方向，或按 lower_threshold / upper_threshold 正常区间判断的异常采样点数量",
        "- 超阈占比：超阈采样点数量 / 有效采样点数量",
        "- 最大连续超阈：同一指标连续处于超阈状态的最长时间",
        "- 标准差 / IQR / 变异系数：用于观察资源波动性",
        "- 趋势斜率：按时间对指标做线性回归，单位为指标值/分钟",
        "- 资源效率：把 CPU、内存、IO、网络、磁盘等资源消耗按吞吐或并发归一化",
        "- 风险评分：综合超阈、连续超阈、泄漏、失败率、响应时间、FPS 等规则扣分",
        "- 发布门禁：按风险等级、泄漏、超阈、失败请求等规则判断是否通过",
        "- 内存斜率：按时间对内存占用做线性回归，单位 MB/分钟",
    ]
    return "\n".join(lines)


def build_generic_report_markdown(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    group_label: str,
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> str:
    summary = summarize(data, thresholds, metrics, threshold_bounds)
    trends = trend_analysis(data, metrics)
    alert_summary = alert_continuity_summary(data, thresholds, metrics, threshold_bounds)
    segments = alert_segments(data, thresholds, metrics, threshold_bounds)
    conclusions = build_generic_conclusions(data, thresholds, metrics, threshold_bounds)
    lines = [
        "# 通用时序分析报告",
        "",
        "## 分析结论",
        "",
        *[f"- {item}" for item in conclusions],
        "",
        "## 汇总数据",
        "",
        dataframe_to_markdown(summary) if not summary.empty else "无汇总数据。",
        "",
        "## 告警持续性",
        "",
        dataframe_to_markdown(alert_summary.drop(columns=["最大连续超阈(分钟)"], errors="ignore")) if not alert_summary.empty else "当前阈值下无超阈持续性数据。",
        "",
        "## 异常段",
        "",
        dataframe_to_markdown(segments.drop(columns=["指标key", "持续时长(分钟)"], errors="ignore")) if not segments.empty else "当前阈值下无异常段数据。",
        "",
        "## 全指标趋势斜率",
        "",
        dataframe_to_markdown(trends) if not trends.empty else "无趋势斜率数据。",
        "",
        "## 分析口径",
        "",
        f"- {group_label}：由 Excel sheet 名提供。",
        "- 均值 / 峰值 / P90 / P95 / P99：用于观察整体水平、高水位和尾部风险。",
        "- 超阈占比：按单阈值方向，或按 lower_threshold / upper_threshold 正常区间判断的异常采样点数量 / 有效采样点数量。",
        "- 最大连续超阈：同一指标连续处于超阈状态的最长时间。",
        "- 趋势斜率：按时间对指标做线性回归，单位为指标值/分钟。",
    ]
    return "\n".join(lines)


def build_report_detail_tables(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    slope_threshold: float,
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
    config_table: pd.DataFrame | None = None,
) -> dict[str, pd.DataFrame]:
    decisions = scene_decision_analysis(data, thresholds, metrics, slope_threshold, threshold_bounds)
    gate_status, gate_rules = release_gate_analysis(
        decisions,
        max_allowed_level="中风险",
        fail_on_leak=True,
        fail_on_any_alert=False,
        fail_on_errors=False,
    )
    return {
        "决策分析": decisions,
        "发布门禁": gate_rules,
        "稳定性评分": stability_score_analysis(decisions),
        "瓶颈提示": bottleneck_hints(data, thresholds, metrics),
        "汇总数据": summarize(data, thresholds, metrics, threshold_bounds),
        "告警持续性": alert_continuity_summary(data, thresholds, metrics, threshold_bounds),
        "异常段": alert_segments(data, thresholds, metrics, threshold_bounds),
        "内存趋势检测": leak_analysis(data, slope_threshold),
        "全指标趋势斜率": trend_analysis(data, metrics),
        "资源效率": resource_efficiency_analysis(data),
        "配置建议": config_table if config_table is not None else pd.DataFrame(),
        "_发布门禁状态": pd.DataFrame([{"发布门禁默认判断": gate_status}]),
    }


def build_generic_report_detail_tables(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
) -> dict[str, pd.DataFrame]:
    return {
        "汇总数据": summarize(data, thresholds, metrics, threshold_bounds),
        "告警持续性": alert_continuity_summary(data, thresholds, metrics, threshold_bounds),
        "异常段": alert_segments(data, thresholds, metrics, threshold_bounds),
        "全指标趋势斜率": trend_analysis(data, metrics),
    }


def build_report_pdf_summary_markdown(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    slope_threshold: float,
    metrics: list[Metric],
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
    config_table: pd.DataFrame | None = None,
    config_notes: list[str] | None = None,
    top_n: int = 10,
) -> str:
    tables = build_report_detail_tables(data, thresholds, slope_threshold, metrics, threshold_bounds, config_table)
    decisions = tables["决策分析"]
    gate_rules = tables["发布门禁"]
    gate_status = tables["_发布门禁状态"]["发布门禁默认判断"].iloc[0] if not tables["_发布门禁状态"].empty else "数据不足"
    alert_summary = tables["告警持续性"]
    segments = tables["异常段"]
    leaks = tables["内存趋势检测"]
    trends = tables["全指标趋势斜率"]
    stability = tables["稳定性评分"]
    bottlenecks = tables["瓶颈提示"]

    alert_count = _count_alerts(data, thresholds, metrics, threshold_bounds)
    high_risk = int((decisions["风险等级"] == "高风险").sum()) if "风险等级" in decisions else 0
    medium_risk = int((decisions["风险等级"] == "中风险").sum()) if "风险等级" in decisions else 0
    leak_names = []
    if not leaks.empty and "判断" in leaks:
        leak_names = [f"{row['数据源']} / {row['场景']}" for _, row in leaks[leaks["判断"] == "疑似泄漏"].iterrows()]

    lines = [
        "# 应用资源消耗分析报告",
        "",
        "## 摘要结论",
        "",
        f"- 发布门禁默认判断：{gate_status}。",
        f"- 当前阈值下超阈采样点：{alert_count} 个。",
        f"- 场景风险概况：高风险 {high_risk} 个，中风险 {medium_risk} 个。",
        f"- 内存趋势检测：{'发现疑似泄漏：' + '、'.join(leak_names) if leak_names else '未发现明显泄漏场景'}。",
        f"- PDF 仅保留摘要和 Top {top_n} 明细；完整明细请查看 ZIP 内的 HTML、Markdown 或 report_details.xlsx。",
        "",
        "## 发布门禁",
        "",
        dataframe_to_markdown(gate_rules) if not gate_rules.empty else "无发布门禁数据。",
        "",
        "## 场景风险 Top",
        "",
        dataframe_to_markdown(
            _select_existing_columns(
                _top_rows(decisions, "风险评分", top_n),
                ["数据源", "场景", "风险评分", "风险等级", "最大超阈占比(%)", "最长连续超阈", "内存趋势", "主要原因"],
            )
        )
        if not decisions.empty
        else "无场景风险评级数据。",
        "",
        "## 稳定性评分 Top",
        "",
        dataframe_to_markdown(_top_rows(stability, "稳定性评分", top_n, ascending=True)) if not stability.empty else "无稳定性评分数据。",
        "",
        "## 关键瓶颈提示",
        "",
        dataframe_to_markdown(_top_rows(bottlenecks, None, top_n)) if not bottlenecks.empty else "无瓶颈提示数据。",
        "",
        "## 告警持续性 Top",
        "",
        dataframe_to_markdown(_top_rows(alert_summary, "超阈占比(%)", top_n).drop(columns=["最大连续超阈(分钟)"], errors="ignore"))
        if not alert_summary.empty
        else "当前阈值下无超阈持续性数据。",
        "",
        "## 异常段 Top",
        "",
        dataframe_to_markdown(_top_rows(segments, "持续时长(分钟)", top_n).drop(columns=["指标key", "持续时长(分钟)"], errors="ignore")) if not segments.empty else "当前阈值下无异常段数据。",
        "",
        "## 内存趋势检测",
        "",
        dataframe_to_markdown(_top_rows(leaks, "内存斜率(MB/分钟)", top_n)) if not leaks.empty else "无内存趋势数据。",
        "",
        "## 趋势斜率 Top",
        "",
        dataframe_to_markdown(_top_rows(trends.assign(_abs_slope=trends["趋势斜率/分钟"].abs()) if "趋势斜率/分钟" in trends else trends, "_abs_slope", top_n).drop(columns=["_abs_slope"], errors="ignore"))
        if not trends.empty
        else "无趋势斜率数据。",
        "",
        "## 配置建议",
        "",
        dataframe_to_markdown(config_table) if config_table is not None and not config_table.empty else "未生成配置建议。",
        "",
        *[f"- {item}" for item in (config_notes or [])[:5]],
    ]
    return "\n".join(lines)


def build_generic_report_pdf_summary_markdown(
    data: pd.DataFrame,
    thresholds: dict[str, float],
    metrics: list[Metric],
    group_label: str,
    threshold_bounds: dict[str, tuple[float | None, float | None]] | None = None,
    top_n: int = 10,
) -> str:
    tables = build_generic_report_detail_tables(data, thresholds, metrics, threshold_bounds)
    summary = tables["汇总数据"]
    alert_summary = tables["告警持续性"]
    segments = tables["异常段"]
    trends = tables["全指标趋势斜率"]
    conclusions = build_generic_conclusions(data, thresholds, metrics, threshold_bounds)
    lines = [
        "# 通用时序分析报告",
        "",
        "## 摘要结论",
        "",
        *[f"- {item}" for item in conclusions[:5]],
        f"- PDF 仅保留摘要和 Top {top_n} 明细；完整明细请查看 ZIP 内的 HTML、Markdown 或 report_details.xlsx。",
        "",
        "## 汇总数据摘要",
        "",
        dataframe_to_markdown(_top_rows(summary, None, top_n)) if not summary.empty else "无汇总数据。",
        "",
        "## 告警持续性 Top",
        "",
        dataframe_to_markdown(_top_rows(alert_summary, "超阈占比(%)", top_n).drop(columns=["最大连续超阈(分钟)"], errors="ignore"))
        if not alert_summary.empty
        else "当前阈值下无超阈持续性数据。",
        "",
        "## 异常段 Top",
        "",
        dataframe_to_markdown(_top_rows(segments, "持续时长(分钟)", top_n).drop(columns=["指标key", "持续时长(分钟)"], errors="ignore")) if not segments.empty else "当前阈值下无异常段数据。",
        "",
        "## 趋势斜率 Top",
        "",
        dataframe_to_markdown(_top_rows(trends.assign(_abs_slope=trends["趋势斜率/分钟"].abs()) if "趋势斜率/分钟" in trends else trends, "_abs_slope", top_n).drop(columns=["_abs_slope"], errors="ignore"))
        if not trends.empty
        else "无趋势斜率数据。",
        "",
        "## 分析口径",
        "",
        f"- {group_label}：由 Excel sheet 名提供。",
        "- PDF 报告面向阅读，仅展示核心摘要；完整数据保存在明细 Excel 中。",
    ]
    return "\n".join(lines)


def build_report_details_xlsx_bytes(tables: dict[str, pd.DataFrame]) -> bytes:
    output = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        wrote_sheet = False
        for name, table in tables.items():
            if name.startswith("_") or table is None or table.empty:
                continue
            sheet_name = _safe_sheet_name(name, used)
            table.to_excel(writer, sheet_name=sheet_name, index=False)
            wrote_sheet = True
        if not wrote_sheet:
            pd.DataFrame([{"说明": "当前筛选条件下无可导出的明细表。"}]).to_excel(writer, sheet_name="说明", index=False)
    return output.getvalue()


def build_report_html(markdown_text: str, title: str | None = None) -> str:
    blocks = _markdown_blocks(markdown_text)
    report_title = title or next((str(block["text"]) for block in blocks if block["type"] == "heading" and block["level"] == 1), "分析报告")
    body_parts = [
        '<header class="report-header">',
        f"<h1>{html.escape(report_title)}</h1>",
        f'<p class="meta">导出时间：{html.escape(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))}</p>',
        "</header>",
    ]
    skip_first_h1 = True
    for block in blocks:
        block_type = block["type"]
        if block_type == "heading":
            level = int(block["level"])
            text = html.escape(str(block["text"]))
            if level == 1 and skip_first_h1:
                skip_first_h1 = False
                continue
            body_parts.append(f"<h{level}>{text}</h{level}>")
        elif block_type == "list":
            body_parts.append("<ul>")
            for item in block["items"]:
                body_parts.append(f"<li>{html.escape(str(item))}</li>")
            body_parts.append("</ul>")
        elif block_type == "table":
            headers = [html.escape(str(header)) for header in block["headers"]]
            rows = block["rows"]
            body_parts.append("<div class=\"table-wrap\"><table>")
            body_parts.append("<thead><tr>" + "".join(f"<th>{header}</th>" for header in headers) + "</tr></thead>")
            body_parts.append("<tbody>")
            for row in rows:
                cells = [html.escape(str(cell)) for cell in row]
                body_parts.append("<tr>" + "".join(f"<td>{cell}</td>" for cell in cells) + "</tr>")
            body_parts.append("</tbody></table></div>")
        else:
            text = html.escape(str(block["text"]))
            body_parts.append(f"<p>{text}</p>")
    body = "\n".join(body_parts)
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(report_title)}</title>
  <style>
    :root {{ color-scheme: light; }}
    body {{ font-family: Arial, "Microsoft YaHei", sans-serif; line-height: 1.62; margin: 0; color: #1f2933; background: #f5f7fa; }}
    .report {{ max-width: 1120px; margin: 0 auto; padding: 32px 36px 48px; background: #fff; min-height: 100vh; }}
    .report-header {{ border-bottom: 2px solid #d8dee9; margin-bottom: 28px; padding-bottom: 18px; }}
    h1 {{ margin: 0 0 10px; font-size: 30px; color: #111827; }}
    h2 {{ margin-top: 30px; border-left: 4px solid #2563eb; padding-left: 10px; font-size: 21px; color: #111827; }}
    p {{ margin: 10px 0; }}
    ul {{ margin: 10px 0 16px 22px; padding: 0; }}
    li {{ margin: 5px 0; }}
    .meta {{ color: #667085; font-size: 13px; }}
    .table-wrap {{ overflow-x: auto; margin: 12px 0 22px; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
    th, td {{ border: 1px solid #d0d7de; padding: 7px 8px; text-align: left; vertical-align: top; }}
    th {{ background: #eef2f7; color: #111827; font-weight: 700; }}
    tr:nth-child(even) td {{ background: #fafbfc; }}
    code {{ background: #f3f4f6; padding: 2px 4px; border-radius: 3px; }}
    @media print {{
      body {{ background: #fff; }}
      .report {{ padding: 0; max-width: none; }}
      .table-wrap {{ overflow: visible; }}
      h2 {{ break-after: avoid; }}
      table {{ break-inside: auto; }}
      tr {{ break-inside: avoid; break-after: auto; }}
    }}
  </style>
</head>
<body><main class="report">{body}</main></body>
</html>
"""


def build_report_docx_bytes(markdown_text: str) -> bytes:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise RuntimeError("未安装 python-docx，无法生成 Word 报告。请先执行 pip install python-docx。") from exc

    document = Document()
    for block in _markdown_blocks(markdown_text):
        block_type = block["type"]
        if block_type == "heading":
            document.add_heading(str(block["text"]), level=int(block["level"]))
        elif block_type == "list":
            for item in block["items"]:
                document.add_paragraph(str(item), style="List Bullet")
        elif block_type == "table":
            headers = [str(header) for header in block["headers"]]
            rows = [[str(cell) for cell in row] for row in block["rows"]]
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for index, header in enumerate(headers):
                table.rows[0].cells[index].text = header
            for row in rows:
                cells = table.add_row().cells
                for index, cell in enumerate(row[: len(headers)]):
                    cells[index].text = cell
        else:
            document.add_paragraph(str(block["text"]))

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_report_pdf_bytes(markdown_text: str, title: str | None = None, images: list[tuple[str, bytes]] | None = None) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Image as ReportImage
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ModuleNotFoundError as exc:
        raise RuntimeError("未安装 reportlab，无法生成 PDF 报告。请先执行 pip install reportlab。") from exc

    font_name = "STSong-Light"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:
        font_name = "Helvetica"

    styles = getSampleStyleSheet()
    base_style = ParagraphStyle(
        "ReportBase",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=13,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=base_style,
        fontSize=18,
        leading=24,
        spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "ReportHeading",
        parent=base_style,
        fontSize=13,
        leading=18,
        spaceBefore=12,
        spaceAfter=6,
        textColor=colors.HexColor("#111827"),
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=base_style,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#667085"),
        spaceAfter=10,
    )
    list_style = ParagraphStyle(
        "ReportList",
        parent=base_style,
        leftIndent=10,
        firstLineIndent=-7,
        spaceAfter=3,
    )
    table_cell_style = ParagraphStyle(
        "ReportTableCell",
        parent=base_style,
        fontSize=6,
        leading=8,
        wordWrap="CJK",
    )
    table_header_style = ParagraphStyle(
        "ReportTableHeader",
        parent=table_cell_style,
        textColor=colors.HexColor("#111827"),
    )

    def paragraph(text: object, style: ParagraphStyle = base_style) -> Paragraph:
        return Paragraph(html.escape(str(text)).replace("\n", "<br/>"), style)

    blocks = _markdown_blocks(markdown_text)
    report_title = title or next((str(block["text"]) for block in blocks if block["type"] == "heading" and block["level"] == 1), "分析报告")
    page_size = landscape(A4)
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=report_title,
    )
    available_width = page_size[0] - doc.leftMargin - doc.rightMargin
    story = [
        paragraph(report_title, title_style),
        paragraph(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", meta_style),
    ]
    if images:
        story.append(paragraph("关键图表", heading_style))
        image_width = available_width
        image_height = image_width * 0.52
        for image_title, image_bytes in images:
            story.append(paragraph(image_title, meta_style))
            story.append(ReportImage(io.BytesIO(image_bytes), width=image_width, height=image_height))
            story.append(Spacer(1, 8))

    skip_first_h1 = True
    for block in blocks:
        block_type = block["type"]
        if block_type == "heading":
            level = int(block["level"])
            if level == 1 and skip_first_h1:
                skip_first_h1 = False
                continue
            story.append(paragraph(block["text"], heading_style if level <= 2 else base_style))
        elif block_type == "list":
            for item in block["items"]:
                story.append(paragraph(f"- {item}", list_style))
        elif block_type == "table":
            raw_headers = [str(header) for header in block["headers"]]
            raw_rows = [[str(cell) for cell in row] for row in block["rows"]]
            max_columns_per_table = 8
            for start in range(0, len(raw_headers), max_columns_per_table):
                end = start + max_columns_per_table
                header_slice = raw_headers[start:end]
                row_slices = []
                for row in raw_rows:
                    row_slice = row[start:end]
                    if len(row_slice) < len(header_slice):
                        row_slice.extend([""] * (len(header_slice) - len(row_slice)))
                    row_slices.append(row_slice)
                if start > 0:
                    story.append(paragraph(f"续表：第 {start + 1}-{min(end, len(raw_headers))} 列", meta_style))
                headers = [paragraph(header, table_header_style) for header in header_slice]
                rows = [[paragraph(cell, table_cell_style) for cell in row] for row in row_slices]
                column_count = max(1, len(headers))
                col_widths = [available_width / column_count] * column_count
                table = Table([headers, *rows], colWidths=col_widths, repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("FONTNAME", (0, 0), (-1, -1), font_name),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
                            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d0d7de")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 3),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                            ("TOPPADDING", (0, 0), (-1, -1), 3),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                        ]
                    )
                )
                story.append(table)
                story.append(Spacer(1, 8))
        else:
            story.append(paragraph(block["text"], base_style))
            story.append(Spacer(1, 4))

    doc.build(story)
    return output.getvalue()


def build_report_zip_bytes(
    markdown_report: str,
    html_report: str,
    docx_report: bytes | None = None,
    pdf_report: bytes | None = None,
    details_xlsx: bytes | None = None,
    base_name: str = "analysis_report",
) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{base_name}.md", markdown_report.encode("utf-8-sig"))
        archive.writestr(f"{base_name}.html", html_report.encode("utf-8"))
        if docx_report is not None:
            archive.writestr(f"{base_name}.docx", docx_report)
        if pdf_report is not None:
            archive.writestr(f"{base_name}.pdf", pdf_report)
        if details_xlsx is not None:
            archive.writestr(f"{base_name}_details.xlsx", details_xlsx)
    return output.getvalue()
